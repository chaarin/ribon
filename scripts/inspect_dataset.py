"""QIT-CEMC 데이터 폴더 구조를 조사해 보고서를 만든다. (Track A, 학교 서버에서 실행)

실행: python -m scripts.inspect_dataset /path/to/QIT-CEMC
결과: docs/qit_cemc_structure.md  (이 파일만 git에 올리면 된다)

원본 파일은 읽기만 하고 수정하지 않는다. 큰 파일도 앞부분만 읽는다.
"""
import argparse
import csv
import io
from collections import Counter, defaultdict
from pathlib import Path

from src.config.settings import ROOT_DIR

OUTPUT_PATH = ROOT_DIR / "docs" / "qit_cemc_structure.md"
TEXT_EXT = {".csv", ".txt", ".dat", ".tsv"}
SAMPLES_PER_EXT = 3
HEAD_LINES = 6


def human(n: float) -> str:
    for unit in ("B", "KB", "MB", "GB", "TB"):
        if n < 1024:
            return f"{n:.1f}{unit}"
        n /= 1024
    return f"{n:.1f}PB"


def head_text(path: Path) -> str:
    with open(path, "rb") as f:
        raw = f.read(64 * 1024)
    for enc in ("utf-8-sig", "cp949", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    lines = text.splitlines()[:HEAD_LINES]
    out = "\n".join(line[:300] for line in lines)
    try:
        dialect = csv.Sniffer().sniff("\n".join(lines[:3]))
        n_cols = len(next(csv.reader(io.StringIO(lines[0]), dialect)))
        out += f"\n\n(구분자 {dialect.delimiter!r}, 첫 줄 열 개수 {n_cols})"
    except Exception:
        pass
    return out


def describe_mat(path: Path) -> str:
    try:
        import scipy.io

        info = scipy.io.whosmat(path)
        return "\n".join(f"{name}: shape={shape}, type={dtype}" for name, shape, dtype in info)
    except NotImplementedError:
        pass  # MATLAB v7.3 = HDF5
    except ImportError:
        return "(scipy 미설치: pip install scipy)"
    return describe_h5(path)


def describe_h5(path: Path) -> str:
    try:
        import h5py
    except ImportError:
        return "(h5py 미설치: pip install h5py)"
    lines = []
    with h5py.File(path, "r") as f:
        f.visititems(
            lambda name, obj: lines.append(f"{name}: shape={obj.shape}, dtype={obj.dtype}")
            if isinstance(obj, h5py.Dataset)
            else None
        )
    return "\n".join(lines[:50])


def describe_xlsx(path: Path) -> str:
    try:
        import openpyxl
    except ImportError:
        return "(openpyxl 미설치: pip install openpyxl)"
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        rows = [r for _, r in zip(range(HEAD_LINES), ws.iter_rows(values_only=True))]
        lines.append(f"## 시트 {ws.title} (max_row={ws.max_row}, max_col={ws.max_column})")
        lines += [" | ".join("" if v is None else str(v) for v in r) for r in rows]
    return "\n".join(lines)


def describe_tdms(path: Path) -> str:
    try:
        from nptdms import TdmsFile
    except ImportError:
        return "(npTDMS 미설치: pip install npTDMS)"
    with TdmsFile.open(path) as f:
        return "\n".join(
            f"{g.name}/{c.name}: len={len(c)}, props={dict(list(c.properties.items())[:5])}"
            for g in f.groups()
            for c in g.channels()
        )


DESCRIBERS = {".mat": describe_mat, ".h5": describe_h5, ".hdf5": describe_h5, ".xlsx": describe_xlsx, ".tdms": describe_tdms}


def describe(path: Path) -> str:
    ext = path.suffix.lower()
    try:
        if ext in TEXT_EXT:
            return head_text(path)
        if ext in DESCRIBERS:
            return DESCRIBERS[ext](path)
        return "(형식 미지원: 확장자를 알려주세요)"
    except Exception as e:  # 한 파일 실패가 전체 조사를 막지 않게 한다
        return f"(읽기 실패: {type(e).__name__}: {e})"


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("root", type=Path)
    args = parser.parse_args()
    root: Path = args.root.resolve()

    files = [p for p in root.rglob("*") if p.is_file()]
    by_ext: dict[str, list[Path]] = defaultdict(list)
    dir_stats: dict[Path, list[int]] = defaultdict(lambda: [0, 0])
    for p in files:
        by_ext[p.suffix.lower() or "(없음)"].append(p)
        rel_dir = p.parent.relative_to(root)
        dir_stats[rel_dir][0] += 1
        dir_stats[rel_dir][1] += p.stat().st_size

    out = ["# QIT-CEMC 데이터 구조 보고서\n", f"- 경로: `{root}`", f"- 파일 {len(files)}개, 전체 {human(sum(s for _, s in dir_stats.values()))}\n"]

    out.append("## 확장자별\n\n| 확장자 | 개수 | 용량 |\n|---|---|---|")
    for ext, ps in sorted(by_ext.items(), key=lambda kv: -len(kv[1])):
        out.append(f"| {ext} | {len(ps)} | {human(sum(p.stat().st_size for p in ps))} |")

    out.append("\n## 폴더별 (상위 40개)\n\n| 폴더 | 파일 수 | 용량 |\n|---|---|---|")
    for d, (n, size) in sorted(dir_stats.items(), key=lambda kv: str(kv[0]))[:40]:
        out.append(f"| `{d}` | {n} | {human(size)} |")

    out.append("\n## 파일 이름 예시 (폴더별 앞 5개)\n")
    shown = Counter()
    for p in sorted(files):
        rel = p.relative_to(root)
        if shown[rel.parent] < 5:
            out.append(f"- `{rel}` ({human(p.stat().st_size)})")
            shown[rel.parent] += 1

    out.append("\n## 파일 내용 샘플\n")
    for ext, ps in sorted(by_ext.items()):
        for p in sorted(ps, key=lambda x: x.stat().st_size)[:SAMPLES_PER_EXT]:
            out.append(f"### `{p.relative_to(root)}` ({human(p.stat().st_size)})\n\n```\n{describe(p)}\n```\n")

    OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_PATH.write_text("\n".join(out), encoding="utf-8")
    print(f"보고서 저장: {OUTPUT_PATH}")


if __name__ == "__main__":
    main()
